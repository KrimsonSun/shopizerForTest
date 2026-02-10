#!/usr/bin/env python3
from docx import Document

doc = Document('test-reports/Assignment2_Test_Report.docx')

print('='*80)
print('CURRENT DOCUMENT ANALYSIS')
print('='*80)

# Count images
image_count = 0
image_paragraphs = set()

for i, para in enumerate(doc.paragraphs):
    for run in para.runs:
        if run._element.xpath('.//pic:pic'):
            image_count += 1
            image_paragraphs.add(i)

print(f'\nTotal embedded images: {image_count}')
print(f'Total paragraphs with images: {len(image_paragraphs)}')

# Find screenshot placeholders
print('\n' + '='*80)
print('SCREENSHOT PLACEHOLDERS STATUS')
print('='*80)

screenshot_info = []
for i, para in enumerate(doc.paragraphs):
    text = para.text.strip()
    if '📷' in text or 'SCREENSHOT' in text.upper():
        # Check if image is nearby (within 2 paragraphs)
        has_nearby_image = any(j in image_paragraphs for j in range(max(0, i-2), min(len(doc.paragraphs), i+3)))
        screenshot_info.append({
            'index': i,
            'text': text[:80],
            'has_image': has_nearby_image
        })

for idx, info in enumerate(screenshot_info, 1):
    status = '✅' if info['has_image'] else '❌'
    print(f"{idx:2}. {status} {info['text']}")

print(f"\n✅ Inserted: {sum(1 for s in screenshot_info if s['has_image'])}")
print(f"❌ Missing: {sum(1 for s in screenshot_info if not s['has_image'])}")
